#!/usr/bin/env python3
"""Generate NL Swiggy Instamart 10-slide deck (min 14pt, colorblind-safe)."""

from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx", "-q"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Colorblind-safe palette (blue / orange / navy / slate — no red-green)
NAVY = RGBColor(0x1A, 0x1A, 0x2E)
ORANGE = RGBColor(0xFC, 0x80, 0x19)
BLUE = RGBColor(0x2E, 0x5A, 0xAC)
SLATE = RGBColor(0x33, 0x41, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
SAND = RGBColor(0xFF, 0xF7, 0xED)

W, H = Inches(13.333), Inches(7.5)  # 16:9


def set_run(run, size=14, bold=False, color=SLATE):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"


def add_bg(slide, color=LIGHT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_bar(slide, color=ORANGE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def textbox(slide, left, top, width, height, text, size=14, bold=False, color=SLATE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def bullets(slide, left, top, width, height, lines, size=14, color=SLATE):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = ("• " if not line.startswith("•") else "") + line.lstrip("• ").strip()
        set_run(run, size=size, color=color)
    return box


def new_slide(prs):
    slide = prs.slide_layouts[6]  # blank
    s = prs.slides.add_slide(slide)
    add_bg(s)
    add_bar(s)
    return s


def main():
    out_dir = Path(__file__).resolve().parents[2] / "deliverables"
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Title
    s = new_slide(prs)
    add_bg(s, NAVY)
    add_bar(s, ORANGE)
    textbox(s, Inches(0.7), Inches(1.8), Inches(12), Inches(1),
            "Swiggy Instamart · Growth", size=18, bold=True, color=ORANGE)
    textbox(s, Inches(0.7), Inches(2.5), Inches(12), Inches(2),
            "MACs buy from the same aisles every month —\nhabit, not awareness, is the growth leak",
            size=32, bold=True, color=WHITE)
    textbox(s, Inches(0.7), Inches(5.2), Inches(12), Inches(1),
            "AI discovery engine + validated research + deployed Basket-Gap Agent",
            size=16, color=WHITE)

    # 2 Goal
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.2),
            "North-star: raise % of MACs who buy ≥1 new category each month",
            size=26, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(2), Inches(11.5), Inches(4.5), [
        "Examples: groceries → pet · snacks → personal care · household → baby",
        "Today’s behavior is replenishment-first: same search queries, same SKUs",
        "Winning means unlocking adjacency categories without breaking the 10-min job",
        "Chosen product: Swiggy Instamart (standalone app + ecosystem)",
    ], size=16)

    # 3 Engine findings
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.2),
            "AI engine: reorder habit + search tunnel + first-buy risk dominate",
            size=24, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(5), [
        "Corpus: Play Store + App Store (Instamart/Swiggy/Blinkit/Zepto/BigBasket) + Reddit RSS",
        "Top themes: habit autopilot, search-first navigation, quality anxiety, time pressure",
        "Users discover via typed search & past orders — not category browse",
        "Every insight carries quote IDs; validation scorecard on /workflow",
        "Live demo: paste a review → Groq classification in the workflow UI",
    ], size=15)

    # 4 Workflow 1-slider (REQUIRED)
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(1),
            "How the review-analysis workflow works (testable live)",
            size=24, bold=True, color=NAVY)
    steps = [
        ("1. Ingest", "Play/App/Reddit scrapers → JSONL"),
        ("2. Normalize", "Dedupe · language · substance filter"),
        ("3. Classify", "Groq gpt-oss → blockers & segments"),
        ("4. Theme", "Cluster + name themes with evidence"),
        ("5. Validate", "Gold set · κ · triangulation · citations"),
    ]
    for i, (title, body) in enumerate(steps):
        left = Inches(0.45 + i * 2.5)
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.0), Inches(2.3), Inches(3.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = BLUE
        textbox(s, left + Inches(0.12), Inches(2.2), Inches(2.05), Inches(1),
                title, size=16, bold=True, color=ORANGE)
        textbox(s, left + Inches(0.12), Inches(3.1), Inches(2.05), Inches(1.8),
                body, size=14, color=SLATE)
    textbox(s, Inches(0.6), Inches(5.6), Inches(12), Inches(1),
            "Link: /workflow  ·  Re-run: npm run scrape:all && npm run pipeline && npm run pipeline:validate",
            size=14, color=BLUE)

    # 5 Research
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.2),
            "6 interviews validate habit & risk — and challenge “price is #1”",
            size=24, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(5), [
        "Segment: replenishment-heavy users who buy pet/baby/care/pharmacy elsewhere",
        "Validated: reorder autopilot, search tunnel, first-purchase risk, ops trust multiplier",
        "Challenged: price opacity as primary blocker for adjacency (trust/size/replace win)",
        "Challenged: UI alone — users also frame Instamart as a “kitchen-only” mental aisle",
        "Artefact: /research/research-pack.md (screener, guide, transcripts, matrix)",
    ], size=15)

    # 6 Problem
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.4),
            "Problem: first buys in new categories feel costly under a kitchen-only habit loop",
            size=24, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(2.0), Inches(11.5), Inches(4.5), [
        "Target: urban MACs (25–40) with weekly grocery/snack reorders + offline adjacency spend",
        "Root cause: search-led habits + first-purchase risk + mental model silo",
        "Workarounds: Blinkit/Zepto multi-home, chemist, pet store, supermarket care aisle",
        "User value: one less app/store switch with trusted starter baskets",
        "Business value: ↑ % MACs with new category / month → AOV, retention, category depth",
    ], size=15)

    # 7 Solution
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.2),
            "MVP: Basket-Gap Agent — trigger-timed, de-risked 3-SKU first baskets",
            size=24, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(5), [
        "Infer household context from reorder graph (pet, toddler, gym, allergies)",
        "Detect categories the household almost certainly buys elsewhere",
        "Generate 3 SKUs with plain-language reason + price anchor + risk reversal",
        "Time the nudge to a trigger (restock cycle, season, life stage) — not a generic carousel",
        "Deployed at /agent with seeded personas + live Groq chat",
    ], size=15)

    # 8 Demo
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.2),
            "Production MVP: pick a persona, see the gap, chat the agent, add a trial SKU",
            size=24, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(5), [
        "Personas: Pet parent · Snack repeater · Family/toddler · WFH cook",
        "Mobile UI mirrors Instamart “For you · New for your home” placement",
        "Agent explains why these 3 and how replace promises remove fear",
        "Falls back to cached replies if Groq rate-limits during evaluation",
        "Link: /agent",
    ], size=15)

    # 9 Impact
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.2),
            "Success metric: new-category MAC rate — leading indicators on trial & repeat",
            size=22, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(5), [
        "Primary: % MACs with ≥1 order in a category new-to-them in the last 30 days",
        "Leading: shortlist CTR · add-to-cart · first-order completion in gap category",
        "Quality: replace/refund rate on first-timer SKUs (must not rise)",
        "Guardrail: do not hurt grocery replenishment conversion / ETA satisfaction",
        "Rollout: pet + baby cohorts first (highest interview willingness)",
    ], size=15)

    # 10 Links
    s = new_slide(prs)
    textbox(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1.2),
            "Deliverables: workflow, research pack, and deployed agent",
            size=24, bold=True, color=NAVY)
    bullets(s, Inches(0.8), Inches(1.8), Inches(11.5), Inches(5), [
        "Review analysis workflow (live): /workflow",
        "Basket-Gap Agent MVP (production): /agent",
        "Research pack (screener, guide, 6 interviews, matrix): /research/research-pack.md",
        "Problem & synthesis summary: /research",
        "Stack: Next.js · Groq openai/gpt-oss-120b · Vercel · Play/App/Reddit scrapers",
    ], size=15)

    pptx_path = out_dir / "NL Swiggy Instamart.pptx"
    prs.save(str(pptx_path))
    print(f"Wrote {pptx_path}")

    # Try PDF export via LibreOffice if available
    import shutil, subprocess
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)],
            check=False,
        )
        print("PDF export attempted")
    else:
        print("LibreOffice not found — open the PPTX and File → Export → PDF")


if __name__ == "__main__":
    main()
